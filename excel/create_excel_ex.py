import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from tkinter import filedialog, messagebox


def create_excel_with_headers(title, headers, data_rows=None, parent=None):
    """헤더와 데이터를 이용해 엑셀 파일 생성"""
    # 파일 저장 경로 선택
    download_folder = os.path.expanduser("~/Downloads")
    timestamp = datetime.now().strftime("%Y%m%d")  # 일 단위까지만 표시
    default_filename = f"{title}_{timestamp}.xlsx"

    filepath = filedialog.asksaveasfilename(
        initialdir=download_folder,
        initialfile=default_filename,
        defaultextension=".xlsx",
        filetypes=[("Excel 파일", "*.xlsx"), ("모든 파일", "*.*")]
    )

    if not filepath:
        return None  # 사용자가 취소함

    # 엑셀 파일 생성
    success = generate_excel(filepath, title, headers, data_rows)

    if success:
        messagebox.showinfo("성공", f"엑셀 파일이 생성되었습니다.\n{filepath}")
        return filepath

    return None


def generate_excel(filepath, title, headers, data_rows=None):
    """엑셀 파일 생성 및 저장"""
    try:
        # 워크북 생성
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = title

        # 헤더 추가
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)

        # 데이터 행 추가
        if data_rows and isinstance(data_rows, list) and len(data_rows) > 0:
            for row_idx, item in enumerate(data_rows, start=2):
                for col_idx, header in enumerate(headers, start=1):
                    # 해당 키가 있는 경우에만 값 추가
                    if header in item:
                        value = item[header]
                        ws.cell(row=row_idx, column=col_idx, value=value)

        # 데이터 행 수 설정 (데이터 행이 있으면 데이터 행 수, 없으면 기본 10행)
        data_row_count = len(data_rows) if data_rows else 10

        # 스타일 적용
        apply_excel_style(ws, len(headers), data_row_count)

        # 열 너비 자동 조정
        for col_idx, header in enumerate(headers, start=1):
            column_letter = get_column_letter(col_idx)
            ws.column_dimensions[column_letter].width = max(len(str(header)) * 1.5, 12)

        # 저장
        wb.save(filepath)
        return True
    except Exception as e:
        messagebox.showerror("오류", f"엑셀 파일 생성 중 오류가 발생했습니다: {str(e)}")
        return False


def apply_excel_style(worksheet, header_count, data_row_count):
    """엑셀 워크시트에 스타일 적용"""
    # 테두리 스타일 정의
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # 헤더 행 스타일 설정
    header_row = worksheet[1]
    for cell in header_row:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical='center', horizontal='center')
        cell.fill = PatternFill(start_color="FFA0A0A0", end_color="FFA0A0A0", fill_type="solid")
        cell.border = thin_border

    # 필터 추가
    worksheet.auto_filter.ref = f"A1:{get_column_letter(header_count)}1"

    # 데이터 행에 테두리 스타일 적용
    for row in range(2, data_row_count + 2):
        for col in range(1, header_count + 1):
            cell = worksheet.cell(row=row, column=col)
            cell.border = thin_border